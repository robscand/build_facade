from PyQt5 import QtCore, QtGui, QtWidgets
import mainwindow, inproj, delproj, openproj, netgraph, mkplan, printlist, conndb, chmat, addmat, addwmat

from sqlalchemy import Column, Integer, String, create_engine, or_
from sqlalchemy.ext.declarative import declarative_base
from sqlalchemy.orm import sessionmaker


import cx_Oracle
import datetime

hostname = 'localhost'
port = 1521
dbName = 'xe'

user = 'tripartio'  # tripartio << должен вводиться пользователем
password = 'q1w2r3t4'  # q1w2r3t4 << должен вводиться пользователем

sid = cx_Oracle.makedsn(hostname, port, sid=dbName)

connect_str = 'oracle+cx_oracle://{user}:{password}@{sid}'.format(user=user,
                                                                  password=password,
                                                                  sid=sid)
engine = create_engine(connect_str, echo=False)
Base = declarative_base(engine)


class DescWrk(Base):
    __tablename__ = 'DESC_WRK'
    __table_args__ = {'autoload': True}


class RuleEmp(Base):
    __tablename__ = 'FK_RB_WRK_RB_EMP'
    __table_args__ = {'autoload': True}


class Hours(Base):
    __tablename__ = 'HRS_REQ'
    __table_args__ = {'autoload': True}


class Period(Base):
    __tablename__ = 'PERIOD'
    __table_args__ = {'autoload': True}


class Plan(Base):
    __tablename__ = 'PLAN'
    __table_args__ = {'autoload': True}


class Projs(Base):
    __tablename__ = 'PROJS'
    __table_args__ = {'autoload': True}


class RefBookCompanies(Base):
    __tablename__ = 'RB_COMP'
    __table_args__ = {'autoload': True}


class RefBookEmp(Base):
    __tablename__ = 'RB_EMP'
    __table_args__ = {'autoload': True}


class RefBookMat(Base):
    __tablename__ = 'RB_MAT'
    __table_args__ = {'autoload': True}


class RefBookW(Base):
    __tablename__ = 'RB_WRK'
    __table_args__ = {'autoload': True}


class RuleMat(Base):
    __tablename__ = 'RULE_MAT'
    __table_args__ = {'autoload': True}


class RuleTrip(Base):
    __tablename__ = 'RULE_TRIP'
    __table_args__ = {'autoload': True}


class TaskTerm(Base):
    __tablename__ = 'TASK_TERM'
    __table_args__ = {'autoload': True}


class Unit(Base):
    __tablename__ = 'UNIT'
    __table_args__ = {'autoload': True}


class Week(Base):
    __tablename__ = 'WEEK_REQ'
    __table_args__ = {'autoload': True}


class AddWMat(QtWidgets.QDialog):
    def __init__(self, works, mats, parent=None):
        QtWidgets.QDialog.__init__(self, parent)
        self.ui = addwmat.Ui_Dialog()
        self.ui.setupUi(self)

        self.buttonBox = QtWidgets.QDialogButtonBox(QtWidgets.QDialogButtonBox.Ok | QtWidgets.QDialogButtonBox.Cancel)
        self.buttonBox.setOrientation(QtCore.Qt.Horizontal)
        self.buttonBox.setObjectName("buttonBox")
        self.ui.gridLayout.addWidget(self.buttonBox, 3, 0, 1, 2)

        self.buttonBox.accepted.connect(self.accept)
        self.buttonBox.rejected.connect(self.reject)

        self.ui.comWork.addItems(works)
        self.ui.comMat.addItems(mats)

        self.intValidator = QtGui.QDoubleValidator(0, 1000, 5, self)
        self.ui.inpVal.setValidator(self.intValidator)

    def clear(self):
        self.ui.inpVal.clear()

    def accept(self):
        self.work = self.ui.comWork.itemText(self.ui.comWork.currentIndex())
        self.mat = self.ui.comMat.itemText(self.ui.comMat.currentIndex())
        self.un_val = self.ui.inpVal.text()
        dot = '.'
        self.un_val = self.un_val.split(',')
        self.un_val = float(dot.join(self.un_val))
        QtWidgets.QDialog.accept(self)

    def get_output(self):
        return self.work, self.mat, self.un_val


class AddNewMat(QtWidgets.QDialog):
    def __init__(self, mats, parent=None):
        QtWidgets.QDialog.__init__(self, parent)
        self.ui = addmat.Ui_Dialog()
        self.ui.setupUi(self)

        self.buttonBox = QtWidgets.QDialogButtonBox(QtWidgets.QDialogButtonBox.Ok | QtWidgets.QDialogButtonBox.Cancel)
        self.buttonBox.setOrientation(QtCore.Qt.Horizontal)
        self.buttonBox.setObjectName("buttonBox")
        self.ui.gridLayout.addWidget(self.buttonBox, 3, 1, 1, 1)

        self.buttonBox.accepted.connect(self.accept)
        self.buttonBox.rejected.connect(self.reject)

        self.ui.comboUn.addItems(mats)

        rxName = QtCore.QRegExp("[0-9,а-я,А-Я,\s,№]+")
        self.nameValidator = QtGui.QRegExpValidator(rxName, self)
        self.ui.inpName.setValidator(self.nameValidator)

        self.ui.inpCode.setInputMask('999-999D;_')# ___-____

        self.intValidator = QtGui.QIntValidator(0, 1000, self)
        self.ui.inpUn.setValidator(self.intValidator)

    def clear(self):
        self.ui.inpCode.clear()
        self.ui.inpName.clear()
        self.ui.inpUn.clear()

    def accept(self):
        self.code = self.ui.inpCode.text()
        self.name = self.ui.inpName.text()
        self.un_val = self.ui.inpUn.text()
        self.un_name = self.ui.comboUn.itemText(self.ui.comboUn.currentIndex())
        QtWidgets.QDialog.accept(self)

    def get_output(self):
        return self.code, self.name, self.un_val, self.un_name


class ChMat(QtWidgets.QDialog):
    def __init__(self, parent=None):
        QtWidgets.QDialog.__init__(self, parent)
        self.ui = chmat.Ui_Dialog()
        self.ui.setupUi(self)

        self.buttonBox = QtWidgets.QDialogButtonBox(QtWidgets.QDialogButtonBox.Ok | QtWidgets.QDialogButtonBox.Cancel)
        self.buttonBox.setOrientation(QtCore.Qt.Horizontal)
        self.buttonBox.setObjectName("buttonBox")
        self.ui.verticalLayout.addWidget(self.buttonBox)

        self.buttonBox.accepted.connect(self.accept)
        self.buttonBox.rejected.connect(self.reject)

    def clear(self):
        if self.ui.rButt_new.isChecked():
            self.ui.rButt_new.setChecked(False)
        if self.ui.rButt_old.isChecked():
            self.ui.rButt_old.setChecked(False)

    def accept(self):
        self.out_new = self.ui.rButt_new.isChecked()
        self.out_old = self.ui.rButt_old.isChecked()
        QtWidgets.QDialog.accept(self)

    def get_output(self):
        return self.out_new, self.out_old


class ConnDB(QtWidgets.QDialog):
    def __init__(self, parent=None):
        QtWidgets.QDialog.__init__(self, parent)
        self.ui = conndb.Ui_Dialog()
        self.ui.setupUi(self)

        self.buttonBox = QtWidgets.QDialogButtonBox(QtWidgets.QDialogButtonBox.Ok | QtWidgets.QDialogButtonBox.Cancel)
        self.buttonBox.setOrientation(QtCore.Qt.Horizontal)
        self.buttonBox.setObjectName("buttonBox")
        self.ui.gridLayout.addWidget(self.buttonBox, 1, 0, 1, 2)

        self.buttonBox.accepted.connect(self.accept)
        self.buttonBox.rejected.connect(self.reject)

    def accept(self):
        self.ip = self.ui.lineEdit.text()
        self.name = self.ui.lineEdit_2.text()
        QtWidgets.QDialog.accept(self)

    def get_output(self):
        return self.ip, self.name

class MakePlan(QtWidgets.QDialog):
    def __init__(self, parent=None):
        QtWidgets.QDialog.__init__(self, parent)
        self.ui = mkplan.Ui_Dialog()
        self.ui.setupUi(self)

        rxCode = QtCore.QRegExp("[1-3]{1}[1-2]{1}[1-3]{1}[1-3]{1}")
        self.codeValidator = QtGui.QRegExpValidator(rxCode, self)
        self.ui.inpCode.setValidator(self.codeValidator)

        rxName = QtCore.QRegExp("[0-9,а-я,А-Я,\s,№]+")
        self.nameValidator = QtGui.QRegExpValidator(rxName, self)
        self.ui.inpName.setValidator(self.nameValidator)

        self.buttonBox = QtWidgets.QDialogButtonBox(QtWidgets.QDialogButtonBox.Ok | QtWidgets.QDialogButtonBox.Cancel)
        self.buttonBox.accepted.connect(self.accept)
        self.buttonBox.rejected.connect(self.reject)
        self.buttonBox.setOrientation(QtCore.Qt.Horizontal)
        self.buttonBox.setObjectName("buttonBox")
        self.ui.gridLayout.addWidget(self.buttonBox, 3, 0, 1, 2)

    def clear(self):
        if self.ui.radioButton.isChecked() == True:
            self.ui.radioButton.setChecked(False)
        self.ui.inpCode.clear()
        self.ui.inpName.clear()
        defalt_date = QtCore.QDate(2000, 1, 1)
        self.ui.inpYear.setDate(defalt_date)

    def accept(self):
        try:
            text = self.ui.inpCode.text()
            year_text = self.ui.inpYear.date().toString(QtCore.Qt.ISODate)[:4]
            self.code_txt = text + '-' + year_text
            self.name_txt = self.ui.inpName.text()
            self.optimize = self.ui.radioButton.isChecked()
            QtWidgets.QDialog.accept(self)
        except Exception:
            QtWidgets.QMessageBox.warning(QtWidgets.QMessageBox(), 'Неверный ввод',
                                          'Не все поля заполнены.\nПовторите ввод или отмените его.')

    def get_output(self):
        return self.code_txt, self.name_txt, self.optimize

class InpProj(QtWidgets.QDialog):
    def __init__(self, parent=None):
        QtWidgets.QDialog.__init__(self, parent)
        self.ui = inproj.Ui_dialog()
        self.ui.setupUi(self)

        rxCode = QtCore.QRegExp("[1-3]{1}[1-2]{1}[1-3]{1}[1-3]{1}")
        self.codeValidator = QtGui.QRegExpValidator(rxCode, self)
        self.ui.inpCode.setValidator(self.codeValidator)

        rxName = QtCore.QRegExp("[0-9,а-я,А-Я,\s,№]+")
        self.nameValidator = QtGui.QRegExpValidator(rxName, self)
        self.ui.inpName.setValidator(self.nameValidator)

        self.intValidator = QtGui.QIntValidator(0, 200, self)
        self.ui.inpFullCost.setValidator(self.intValidator)
        self.ui.inpExpectCost.setValidator(self.intValidator)
        self.ui.inpAllVal.setValidator(self.intValidator)
        self.ui.inpOwnVal.setValidator(self.intValidator)
        self.ui.inpExternVal.setValidator(self.intValidator)

        self.longintValidator = QtGui.QIntValidator(0, 2000, self)
        self.ui.inpReqCap.setValidator(self.longintValidator)

        self.ui.statusBox.addItem("задельный")
        self.ui.statusBox.addItem("пусковой")

        self.buttonBox = QtWidgets.QDialogButtonBox(QtWidgets.QDialogButtonBox.Ok | QtWidgets.QDialogButtonBox.Cancel)
        self.buttonBox.accepted.connect(self.accept)
        self.buttonBox.rejected.connect(self.reject)
        self.buttonBox.setOrientation(QtCore.Qt.Horizontal)
        self.buttonBox.setObjectName("buttonBox")
        self.ui.gridLayout.addWidget(self.buttonBox, 11, 0, 1, 4)

    def accept(self):
        try:
            text = self.ui.inpCode.text()
            year_text = self.ui.inpYear.date().toString(QtCore.Qt.ISODate)[:4]
            self.code_txt = text + '-' + year_text
            self.name_txt = self.ui.inpName.text()
            self.status_txt = self.ui.statusBox.itemText(self.ui.statusBox.currentIndex())
            self.full_cost_num = int(self.ui.inpFullCost.text())
            rdate = self.ui.inpStartDate.date()
            self.start_date = datetime.datetime(rdate.year(), rdate.month(), rdate.day(), 0, 0)
            self.expect_cost_num = int(self.ui.inpExpectCost.text())
            self.req_cap_num = int(self.ui.inpReqCap.text())
            if self.ui.radioButton.isChecked() == True:
                self.fin_date = None
            else:
                rdate = self.ui.inpFinDate.date()
                self.fin_date = datetime.datetime(rdate.year(), rdate.month(), rdate.day(), 0, 0)
            self.all_val_num = int(self.ui.inpAllVal.text())
            self.own_val_num = int(self.ui.inpOwnVal.text())
            self.extern_val_num = int(self.ui.inpExternVal.text())
            QtWidgets.QDialog.accept(self)
        except Exception:
            QtWidgets.QMessageBox.warning(QtWidgets.QMessageBox(), 'Неверный ввод',
                                          'Не все поля заполнены.\nПовторите ввод или отмените его.')

    def get_output(self):
        return self.code_txt, self.name_txt, self.status_txt, \
               self.full_cost_num, self.start_date, self.expect_cost_num, \
               self.req_cap_num, self.fin_date, self.all_val_num, \
               self.own_val_num, self.extern_val_num

    def set_data(self, in_code, in_name, in_status, in_full, in_start, in_expect,
        in_req, in_fin, in_all, in_own, in_extern):
        code, year = in_code[:4], in_code[-4:]
        self.ui.inpCode.setText(code)
        ydate = datetime.datetime.strptime(year, "%Y")
        self.ui.inpYear.setDate(ydate)
        self.ui.inpName.setText(in_name)
        self.ui.statusBox.setCurrentText(in_status)
        self.ui.inpFullCost.setText(str(int(in_full)))
        self.ui.inpStartDate.setDate(in_start)
        self.ui.inpExpectCost.setText(str(int(in_expect)))
        self.ui.inpReqCap.setText(str(int(in_req)))
        if in_fin == 1:
            self.ui.radioButton.setChecked(True)
        else:
            self.ui.inpFinDate.setDate(in_fin)
        self.ui.inpAllVal.setText(str(int(in_all)))
        self.ui.inpOwnVal.setText(str(int(in_own)))
        self.ui.inpExternVal.setText(str(int(in_extern)))

    def clear(self):
        if self.ui.radioButton.isChecked() == True:
            self.ui.radioButton.setChecked(False)
        self.ui.inpCode.clear()
        self.ui.inpName.clear()
        defalt_date = QtCore.QDate(2000, 1, 1)
        self.ui.inpYear.setDate(defalt_date)
        self.ui.inpStartDate.setDate(defalt_date)
        self.ui.inpFinDate.setDate(defalt_date)
        self.ui.inpFullCost.clear()
        self.ui.inpExpectCost.clear()
        self.ui.inpReqCap.clear()
        self.ui.inpAllVal.clear()
        self.ui.inpOwnVal.clear()
        self.ui.inpFullCost.clear()
        self.ui.inpExternVal.clear()

class OpenProj(QtWidgets.QDialog):
    def __init__(self, parent=None):
        QtWidgets.QDialog.__init__(self, parent)
        self.ui = openproj.Ui_Dialog()
        self.ui.setupUi(self)

        self.buttonBox = QtWidgets.QDialogButtonBox(QtWidgets.QDialogButtonBox.Ok | QtWidgets.QDialogButtonBox.Cancel)
        self.buttonBox.accepted.connect(self.accept)
        self.buttonBox.rejected.connect(self.reject)
        self.buttonBox.setOrientation(QtCore.Qt.Horizontal)
        self.buttonBox.setObjectName("buttonBox")
        self.ui.gridLayout.addWidget(self.buttonBox, 1, 0, 1, 2)

    def accept(self):
        self._output = self.ui.dateEdit.date().toString(QtCore.Qt.ISODate)[:4]
        QtWidgets.QDialog.accept(self)

    def get_output(self):
        return self._output

class DelProj(QtWidgets.QDialog):
    def __init__(self, parent=None):
        QtWidgets.QDialog.__init__(self, parent)
        self.ui = delproj.Ui_Dialog()
        self.ui.setupUi(self)

        rxCode = QtCore.QRegExp("[1-3]{1}[1-2]{1}[1-3]{1}[1-3]{1}")
        self.codeValidator = QtGui.QRegExpValidator(rxCode, self)
        self.ui.inpCode.setValidator(self.codeValidator)

        rxName = QtCore.QRegExp("[0-9,а-я,А-Я,\s,№]+")
        self.nameValidator = QtGui.QRegExpValidator(rxName, self)
        self.ui.inpName.setValidator(self.nameValidator)

        self.buttonBox = QtWidgets.QDialogButtonBox(QtWidgets.QDialogButtonBox.Ok | QtWidgets.QDialogButtonBox.Cancel)
        self.buttonBox.accepted.connect(self.accept)
        self.buttonBox.rejected.connect(self.reject)
        self.buttonBox.setOrientation(QtCore.Qt.Horizontal)
        self.buttonBox.setObjectName("buttonBox")
        self.ui.gridLayout.addWidget(self.buttonBox, 2, 0, 1, 2)

    def clear(self):
        self.ui.inpCode.clear()
        self.ui.inpName.clear()
        defalt_date = QtCore.QDate(2000, 1, 1)
        self.ui.inpYear.setDate(defalt_date)

    def accept(self):
        try:
            text = self.ui.inpCode.text()
            year_text = self.ui.inpYear.date().toString(QtCore.Qt.ISODate)[:4]
            self.code_txt = text + '-' + year_text
            self.name_txt = self.ui.inpName.text()
            QtWidgets.QDialog.accept(self)
        except Exception:
            QtWidgets.QMessageBox.warning(QtWidgets.QMessageBox(), 'Неверный ввод',
                                          'Не все поля заполнены.\nПовторите ввод или отмените его.')

    def get_output(self):
        return self.code_txt, self.name_txt


class MainWindow(QtWidgets.QMainWindow):
    def __init__(self, session, parent=None):
        QtWidgets.QMainWindow.__init__(self, parent)

        self.aboutWin = None
        self.addWin = None
        self.delWin = None
        self.openProjWin = None
        self.makePlanWin = None
        self.connDB = None
        self.addNewMat = None
        self.addOldMat = None
        self.chMat = None

        self.ui = mainwindow.Ui_MainWindow()
        self.ui.setupUi(self)

        self.ui.tableWidget.setAlternatingRowColors(True)
        self.ui.tableWidget.horizontalHeader().setCascadingSectionResizes(False)
        self.ui.tableWidget.horizontalHeader().setSectionResizeMode(QtWidgets.QHeaderView.ResizeToContents)
        self.ui.tableWidget.horizontalHeader().setSortIndicatorShown(False)
        self.ui.tableWidget.horizontalHeader().setStretchLastSection(True)
        self.ui.tableWidget.verticalHeader().setVisible(False)
        self.ui.tableWidget.verticalHeader().setCascadingSectionResizes(False)
        self.ui.tableWidget.verticalHeader().setStretchLastSection(False)
        self.ui.tableWidget.verticalHeader().setSectionResizeMode(QtWidgets.QHeaderView.ResizeToContents)

        self.ui.openButton.clicked.connect(self.on_clicked_openButton)
        self.ui.about_action.triggered.connect(self.open_about)
        self.ui.add.triggered.connect(self.open_add)
        self.ui.act_del.triggered.connect(self.open_del)
        self.ui.refresh.triggered.connect(self.open_refresh)
        self.ui.act_edit.triggered.connect(self.open_edit)
        self.ui.print.triggered.connect(self.print)
        self.ui.savePDF.triggered.connect(self.printPDF)
        self.ui.saveWord.triggered.connect(self.printWord)
        self.ui.saveExcel.triggered.connect(self.printExcel)
        self.ui.act_db.triggered.connect(self.conn)

    @QtCore.pyqtSlot()
    def on_clicked_openButton(self):
        try:
            print(self.ui.treeWidget.currentItem().text(self.ui.treeWidget.currentColumn()))
            if self.ui.treeWidget.currentItem().text(self.ui.treeWidget.currentColumn()) == 'Справочник трудовых ресурсов':
                self.loadRefBookPers()
            if self.ui.treeWidget.currentItem().text(self.ui.treeWidget.currentColumn()) == 'Справочник материальных ресурсов':
                self.loadMat()
            if self.ui.treeWidget.currentItem().text(self.ui.treeWidget.currentColumn()) == 'График производства СМР на объект':
                self.loadPlan()
            if self.ui.treeWidget.currentItem().text(self.ui.treeWidget.currentColumn()) == 'Программа строительства на год':
                self.loadProjs()
        except Exception:
            QtWidgets.QMessageBox.information(QtWidgets.QMessageBox(), 'Неверный ввод', 'Документ не был выбран.')

    def loadRefBookPers(self):
        data = {'Kitty': ['1', '2', '3', '4'],
                'Cat': ['4', '5', '6', '7'],
                'Meow': ['5', '8', '3', '1'],
                'Purr': ['6', '3', '9', '2']}
        self.ui.tableWidget.setRowCount(4)
        self.ui.tableWidget.setColumnCount(4)

        horHeaders = []
        for n, key in enumerate(sorted(data.keys())):
            horHeaders.append(key)
            for m, item in enumerate(data[key]):
                newitem = QtWidgets.QTableWidgetItem(item)
                self.ui.tableWidget.setItem(m, n, newitem)

        self.ui.tableWidget.setHorizontalHeaderLabels(horHeaders)

    def get_out_data(self, inp_data, optimize, date):
        # Определяем число вершин
        ng = []
        max_vert = inp_data[-1][1]
        min_vert = inp_data[0][0]
        # Так как алгоритм считает вершины от 0, а в сетевом графике они считаются от 1,
        # то добавляем нулевую фиктивную вершину
        g = netgraph.NetGraph(max_vert + 1)
        for el in inp_data:
            # Добавляем на граф все данные из БД
            g.addEdge(el[0], el[1], el[2], el[3])
        # Расчитываем сетевой график до корректировки
        ng = g.calculateNetGraph()
        # Определяем величину критического пути до
        nocrt = g.getCritWay(save_val=True)
        # Определяем все пути графа
        g.buildAllPaths(min_vert, max_vert, False)
        # Если выбрана корректировка
        if optimize:
            # Расчитываем по алгоритму новые параметры графика
            new_ng = g.optimizeNetGraph()
            ocrt = g.getCritWay(save_val=True)
        else:
            # Оставляем старый график
            new_ng = ng
            ocrt = 0

        # Получаем выходные значения
        # даты начала и окончания
        crews, times, start_dates, end_dates = [], [], [], []
        # дата начала стройки
        start = date
        for i in range(len(new_ng)):
            # если работа первая
            if new_ng[i][0] == min_vert:
                start_dates.append(start)
                end_dates.append(start_dates[i] + datetime.timedelta(days=new_ng[i][3]))
                crews.append(new_ng[i][2])
                times.append(new_ng[i][3])
            else:
                max_l = []
                for j in range(i + 1):
                    if new_ng[i][0] == new_ng[j][1]:
                        # сохраняем рание окончания предыдущих работ
                        max_l.append(new_ng[j][5])
                # начальная дата + наибольшее ранее окончание предыдущих (предыдущей) работ
                start_dates.append(start + datetime.timedelta(days=max(max_l)))
                end_dates.append(start_dates[i] + datetime.timedelta(days=new_ng[i][3]))
                crews.append(new_ng[i][2])
                times.append(new_ng[i][3])
        return nocrt, ocrt, crews, times, start_dates, end_dates

    def loadMat(self):
        horHeaders = ['Обоснование', 'Шифр материала', 'Наименование\nматериала', 'Единица\nизмерения',
                      'Норма расхода\nна единицу\nобъема']

        self.ui.tableWidget.setRowCount(0)
        self.ui.tableWidget.setColumnCount(len(horHeaders))
        self.ui.tableWidget.setHorizontalHeaderLabels(horHeaders)

        mat_dct = {}
        query_mat = session.query(RuleMat).all()
        for idx, obj in enumerate(query_mat):
            param_mat = session.query(RefBookMat).filter(RefBookMat.code_mat == obj.rb_mat_code_mat).first()
            unit = session.query(Unit.name_un).filter(Unit.id_unit == param_mat.unit_id_unit).first()
            mat_lst = [param_mat.code_mat, param_mat.name, str(int(param_mat.unit_val)) + str(unit.name_un),
                       obj.mat_unit]
            self.upMatDct(mat_dct, obj.rb_wrk_num_seen, mat_lst)
        for row_num, row in enumerate(mat_dct.items()):
            self.ui.tableWidget.insertRow(row_num)
            self.ui.tableWidget.setItem(row_num, 0, QtWidgets.QTableWidgetItem(row[0]))
            code = ''
            name = ''
            un = ''
            val = ''
            for element in row[1]:
                code += element[0] + '\n'
                name += element[1] + '\n'
                un += element[2] + '\n'
                val += str(element[3]) + '\n'
            self.ui.tableWidget.setItem(row_num, 1, QtWidgets.QTableWidgetItem(code[:-1]))
            self.ui.tableWidget.setItem(row_num, 2, QtWidgets.QTableWidgetItem(name[:-1]))
            self.ui.tableWidget.setItem(row_num, 3, QtWidgets.QTableWidgetItem(un[:-1]))
            self.ui.tableWidget.setItem(row_num, 4, QtWidgets.QTableWidgetItem(val[:-1]))


    def loadPlan(self):
        if not self.makePlanWin:
            self.makePlanWin = MakePlan(self)
        self.makePlanWin.clear()
        if self.makePlanWin.exec_() == MakePlan.Accepted:
            code, name, optimize = self.makePlanWin.get_output()
            try:
                horHeaders = ['Наименование работ', 'Единица измерения', 'Объем работ', 'Срок выполнения работ', \
                              'Дата начала работ', 'Дата окончания работ', 'Исполнитель']
                project = session.query(Projs.start_date).filter(Projs.code_proj == code, Projs.name_proj == name).first()
                query_res = session.query(Plan, DescWrk). \
                    outerjoin(DescWrk, Plan.rb_wrk_num_seen == DescWrk.rb_wrk_num_seen). \
                    filter(Plan.projs_code_proj == code, Plan.projs_name_proj == name). \
                    filter(or_(DescWrk.projs_name_proj == name, DescWrk.projs_name_proj == None)).all()
                inp_data = []
                for res in query_res:
                    if not res[1]:
                        inp_data.append([int(res[0].start_event),
                                         int(res[0].end_event), 0, 0])
                    else:
                        inp_data.append([int(res[0].start_event),
                                         int(res[0].end_event),
                                         int(res[1].num_crew),
                                         int(res[1].wrk_time)])
                inp_data.sort()

                nocrt, ocrt, crews, times, start_dates, end_dates = self.get_out_data(inp_data, optimize, project.start_date)
                if optimize:
                    QtWidgets.QMessageBox.information(QtWidgets.QMessageBox(), 'Расчеты завершены',
                    'Продолжительность сетевого графика до корректировки: ' + str(nocrt) + '\n' +
                    'Продолжительность сетевого графика после корректировки: ' + str(ocrt))

                # Осуществляем запись результатов в БД
                plan = session.query(Plan).filter(Plan.projs_code_proj == code, Plan.projs_name_proj == name)
                for i, wrk in enumerate(plan):
                    wrk.new_crew = crews[i]
                    wrk.new_time = times[i]
                    wrk.start_date = start_dates[i]
                    wrk.end_date = end_dates[i]
                session.commit()

                # Выводим документ на экран
                fict = 'фиктив.'
                query = session.query(Plan).filter(Plan.rb_wrk_num_seen != fict, Plan.projs_code_proj == code,
                                                   Plan.projs_name_proj == name).all()
                self.ui.tableWidget.setRowCount(0)
                self.ui.tableWidget.setColumnCount(len(horHeaders))
                self.ui.tableWidget.setHorizontalHeaderLabels(horHeaders)
                for i, el in enumerate(query):
                    query_val = session.query(DescWrk.wrk_val). \
                        filter(DescWrk.rb_wrk_num_seen == el.rb_wrk_num_seen).first()
                    query_name = session.query(RefBookW.name, RefBookW.unit_val, RefBookW.unit_id_unit). \
                        filter(RefBookW.num_seen == el.rb_wrk_num_seen).first()
                    query_unit = session.query(Unit.name_un).filter(Unit.id_unit == query_name.unit_id_unit).first()
                    query_comp = session.query(RefBookCompanies.companies). \
                        filter(RefBookCompanies.id_comp == el.rb_comp_id_comp).first()
                    self.ui.tableWidget.insertRow(i)
                    self.ui.tableWidget.setItem(i, 0, QtWidgets.QTableWidgetItem(query_name.name))
                    self.ui.tableWidget.setItem(i, 1, QtWidgets.QTableWidgetItem(str(int(query_name.unit_val)) +
                                                                                 str(query_unit.name_un)))
                    self.ui.tableWidget.setItem(i, 2, QtWidgets.QTableWidgetItem(str(query_val.wrk_val)))
                    self.ui.tableWidget.setItem(i, 3, QtWidgets.QTableWidgetItem(str(int(el.new_time))))
                    self.ui.tableWidget.setItem(i, 4, QtWidgets.QTableWidgetItem(el.start_date.strftime('%d.%m.%Y')))
                    self.ui.tableWidget.setItem(i, 5, QtWidgets.QTableWidgetItem(el.end_date.strftime('%d.%m.%Y')))
                    self.ui.tableWidget.setItem(i, 6, QtWidgets.QTableWidgetItem(query_comp.companies))
            except Exception:
                QtWidgets.QMessageBox.warning(QtWidgets.QMessageBox(), 'Cформировать выходной документ не удалось',
                                          'Проверьте код объекта и наименование на уникальность.')

    def loadProjs(self):
        if not self.openProjWin:
            self.openProjWin = OpenProj(self)
        if self.openProjWin.exec_() == OpenProj.Accepted:
            year = self.openProjWin.get_output()

            horHeaders = ['Код объекта', 'Наименование\nобъекта', 'Статус объекта', 'Полная\nсметная стоимость',
                      'Срок начала\nстроительства', 'Ожидаемое выполнение работ\nк началу года, млн.руб',
                      'Вводимая мощность\nв текущем году, м2', 'Срок ввода в\nэксплуатацию', 'Весь объем работ\nна текущий год, млн.руб',
                      'Объем работ\nорганизации, млн.руб', 'Объем работ\nсубподряда, млн.руб']

            result = session.query(Projs).filter(Projs.code_proj.like('%' + year)).all()
            self.ui.tableWidget.setRowCount(0)
            self.ui.tableWidget.setColumnCount(len(horHeaders))
            self.ui.tableWidget.setHorizontalHeaderLabels(horHeaders)
            for row_num, cell in enumerate(result):
                self.ui.tableWidget.insertRow(row_num)
                self.ui.tableWidget.setItem(row_num, 0, QtWidgets.QTableWidgetItem(cell.code_proj))
                self.ui.tableWidget.setItem(row_num, 1, QtWidgets.QTableWidgetItem(cell.name_proj))
                self.ui.tableWidget.setItem(row_num, 2, QtWidgets.QTableWidgetItem(cell.status))
                self.ui.tableWidget.setItem(row_num, 3, QtWidgets.QTableWidgetItem(str(cell.full_cost)))
                self.ui.tableWidget.setItem(row_num, 4, QtWidgets.QTableWidgetItem(cell.start_date.strftime('%d.%m.%Y')))
                self.ui.tableWidget.setItem(row_num, 5, QtWidgets.QTableWidgetItem(str(cell.expect_cost)))
                self.ui.tableWidget.setItem(row_num, 6, QtWidgets.QTableWidgetItem(str(cell.req_cap)))
                if not cell.fin_date:
                    self.ui.tableWidget.setItem(row_num, 7, QtWidgets.QTableWidgetItem(''))
                else:
                    self.ui.tableWidget.setItem(row_num, 7, QtWidgets.QTableWidgetItem(cell.fin_date.strftime('%d.%m.%Y')))
                self.ui.tableWidget.setItem(row_num, 8, QtWidgets.QTableWidgetItem(str(cell.all_val)))
                self.ui.tableWidget.setItem(row_num, 9, QtWidgets.QTableWidgetItem(str(cell.own_val)))
                self.ui.tableWidget.setItem(row_num, 10, QtWidgets.QTableWidgetItem(str(cell.extern_val)))

    def addMat(self):
        if not self.chMat:
            self.chMat = ChMat(self)
        self.chMat.clear()
        if self.chMat.exec_() == ChMat.Accepted:
            nw, od = self.chMat.get_output()
            if nw:
                mats = []
                for name in session.query(Unit.name_un):
                    mats.append(name[0])
                if not self.addNewMat:
                    self.addNewMat = AddNewMat(mats, self)
                self.addNewMat.clear()
                if self.addNewMat.exec_() == AddNewMat.Accepted:
                    code, name, un_val, un_name = self.addNewMat.get_output()
                    id_un = session.query(Unit.id_unit).filter(Unit.name_un == un_name).first()
                    inp_new_mat = RefBookMat(code_mat=code, name=name, unit_val=int(un_val), unit_id_unit=id_un[0])
                    session.add(inp_new_mat)
                    session.commit()
            if od:
                mats = []
                for name in session.query(RefBookMat.name):
                    mats.append(name[0])
                works = []
                for name in session.query(RefBookW.num_seen):
                    works.append(name[0])
                if not self.addOldMat:
                    self.addOldMat = AddWMat(works, mats, self)
                self.addOldMat.clear()
                if self.addOldMat.exec_() == AddWMat.Accepted:
                    work, mat, un_val = self.addOldMat.get_output()
                    id_mat = session.query(RefBookMat.code_mat).filter(RefBookMat.name == mat).first()
                    inp_old_mat = RuleMat(rb_wrk_num_seen=work, rb_mat_code_mat=id_mat[0], mat_unit=un_val)
                    session.add(inp_old_mat)
                    session.commit()

    def addProjs(self):
        if not self.addWin:
            self.addWin = InpProj(self)
        self.addWin.clear()
        if self.addWin.exec_() == InpProj.Accepted:
            code, name, status, full_cost, start_date, expect_cost, \
            req_cap, fin_date, all_val, own_val, extern_val = self.addWin.get_output()
            try:
                new_proj = Projs(code_proj=code, name_proj=name, \
                         status=status, full_cost=full_cost, \
                         start_date=start_date, expect_cost=expect_cost, \
                         unit_id_unit=8, req_cap=req_cap, fin_date=fin_date, all_val=all_val, own_val=own_val,
                         extern_val=extern_val)
                session.add(new_proj)
                session.commit()
            except Exception:
                QtWidgets.QMessageBox.warning(QtWidgets.QMessageBox(), 'Добавление не удалось',
                                              'Проверьте код объекта и наименование на уникальность.')

    @QtCore.pyqtSlot()
    def printExcel(self):
        # Подключаем класс рабочей книги
        from openpyxl import Workbook

        f = QtWidgets.QFileDialog.getSaveFileName(parent=self,
                                                  caption='Сохранение документа',
                                                  directory=QtCore.QDir.currentPath(),
                                                  filter="All (*);;Документ Excel (*.xlsx)")
        filename = f[0]

        cols = self.ui.tableWidget.columnCount()
        headers = [self.ui.tableWidget.horizontalHeaderItem(col).text() for col in range(cols)]
        if not headers:
            QtWidgets.QMessageBox.warning(QtWidgets.QMessageBox(), 'Cохранение не удалось',
                                          'Чтобы сохранить документ, необходимо его открыть.')
            return

        rows = self.ui.tableWidget.rowCount()
        data = []
        for row in range(rows):
            data.append([self.ui.tableWidget.item(row, col).text() for col in range(cols)])

        # Создаем экземпляр рабочей книги
        wb = Workbook()
        # Задаем имя файла, в который запишем книгу
        #doc = self.ui.treeWidget.currentItem().text(self.ui.treeWidget.currentColumn())
        #fname = doc + '.xlsx'
        # Получаем первый лист
        ws = wb.active #ws1 = wb.worksheets[0]
        ws.append(headers)
        for idx in range(1, rows):
            ws.append(data[idx])
        # Сохраняем рабочую книгу
        wb.save(filename=filename)

    @QtCore.pyqtSlot()
    def printWord(self):
        from docx import Document

        f = QtWidgets.QFileDialog.getSaveFileName(parent=self,
                                                  caption='Сохранение документа',
                                                  directory=QtCore.QDir.currentPath(),
                                                  filter="All (*);;Документ Word (*.docx)")
        filename = f[0]

        cols = self.ui.tableWidget.columnCount()
        headers = [self.ui.tableWidget.horizontalHeaderItem(col).text() for col in range(cols)]
        if not headers:
            QtWidgets.QMessageBox.warning(QtWidgets.QMessageBox(), 'Cохранение не удалось',
                                          'Чтобы сохранить документ, необходимо его открыть.')
            return

        rows = self.ui.tableWidget.rowCount()
        data = []
        for row in range(rows):
            data.append([self.ui.tableWidget.item(row, col).text() for col in range(cols)])

        # сохраняем данные в ворде
        document = Document()
        table = document.add_table(rows=1, cols=cols)
        hdr_cells = table.rows[0].cells
        for idx, head in enumerate(headers):
            hdr_cells[idx].text = head
        for row in data:
            row_cells = table.add_row().cells
            for i in range(cols):
                row_cells[i].text = str(row[i])

        #doc = self.ui.treeWidget.currentItem().text(self.ui.treeWidget.currentColumn())
        table.style = 'TableGrid'
        document.save(filename)

    @QtCore.pyqtSlot()
    def print(self):
        cols = self.ui.tableWidget.columnCount()
        headers = [self.ui.tableWidget.horizontalHeaderItem(col).text() for col in range(cols)]
        if not headers:
            QtWidgets.QMessageBox.warning(QtWidgets.QMessageBox(), 'Cохранение не удалось',
                                          'Чтобы сохранить документ, необходимо его открыть.')
            return

        rows = self.ui.tableWidget.rowCount()
        data = []
        for row in range(rows):
            data.append([self.ui.tableWidget.item(row, col).text() for col in range(cols)])

        pl = printlist.PrintList()
        pl.data = data
        pl.headers = headers
        if self.ui.treeWidget.currentItem().text(self.ui.treeWidget.currentColumn()) == 'График производства СМР на объект':
            pl.columnWidths = [205, 85, 60, 85, 85, 85, 125]
        pl.printData()

    @QtCore.pyqtSlot()
    def printPDF(self):
        f = QtWidgets.QFileDialog.getSaveFileName(parent=self,
                                                         caption='Сохранение документа',
                                                         directory=QtCore.QDir.currentPath(),
                                                         filter="All (*);;Документ PDF (*.pdf)")
        filename = f[0]

        cols = self.ui.tableWidget.columnCount()
        headers = [self.ui.tableWidget.horizontalHeaderItem(col).text() for col in range(cols)]
        if not headers:
            QtWidgets.QMessageBox.warning(QtWidgets.QMessageBox(), 'Cохранение не удалось',
                                          'Чтобы сохранить документ, необходимо его открыть.')
            return

        rows = self.ui.tableWidget.rowCount()
        data = []
        for row in range(rows):
            data.append([self.ui.tableWidget.item(row, col).text() for col in range(cols)])

        pl = printlist.PrintList()
        pl.data = data
        pl.headers = headers
        #doc = self.ui.treeWidget.currentItem().text(self.ui.treeWidget.currentColumn())
        #pl.printer.setOutputFileName(doc + '.pdf')
        pl.printer.setOutputFileName(filename)
        #if doc == 'График производства СМР на объекте':
        pl.columnWidths = [210, 90, 65, 90, 90, 90, 130]
        pl.printData()

    def editProjs(self):
        if not self.delWin:
            self.delWin = DelProj(self)
        self.delWin.clear()
        if self.delWin.exec_() == DelProj.Accepted:
            code, name = self.delWin.get_output()
            try:
                editquery = session.query(Projs).filter(Projs.code_proj == code, Projs.name_proj == name).first()
                if not self.addWin:
                    self.addWin = InpProj(self)
                self.addWin.clear()
                if not editquery.fin_date:
                    fin_date = 1
                else:
                    fin_date = editquery.fin_date
                self.addWin.set_data(editquery.code_proj, editquery.name_proj, editquery.status, editquery.full_cost,
                                     editquery.start_date, editquery.expect_cost, editquery.req_cap, fin_date,
                                     editquery.all_val, editquery.own_val, editquery.extern_val)
                if self.addWin.exec_() == InpProj.Accepted:
                    code, name, status, full_cost, start_date, expect_cost, \
                    req_cap, fin_date, all_val, own_val, extern_val = self.addWin.get_output()

                    print(code, name, status, full_cost, start_date, expect_cost, \
                    req_cap, fin_date, all_val, own_val, extern_val)

                    editquery.code_proj=code
                    editquery.name_proj=name
                    editquery.status=status
                    editquery.full_cost=full_cost
                    editquery.start_date=start_date
                    editquery.expect_cost=expect_cost
                    editquery.unit_id_unit=8
                    editquery.req_cap=req_cap
                    editquery.fin_date=fin_date
                    editquery.all_val=all_val
                    editquery.own_val=own_val
                    editquery.extern_val=extern_val

                    session.add(editquery)
                    session.commit()
            except Exception:
                QtWidgets.QMessageBox.warning(QtWidgets.QMessageBox(), 'Редактирование не удалось',
                                              'Проверьте код объекта и наименование на правильность ввода.')

    def delProjs(self):
        if not self.delWin:
            self.delWin = DelProj(self)
        self.delWin.clear()
        if self.delWin.exec_() == DelProj.Accepted:
            code, name = self.delWin.get_output()
            try:
                delquery = session.query(Projs).filter(Projs.code_proj == code, Projs.name_proj == name).first()
                session.delete(delquery)
                session.commit()
            except Exception:
                QtWidgets.QMessageBox.warning(QtWidgets.QMessageBox(), 'Удаление не удалось',
                                          'Проверьте код объекта и наименование на правильность ввода.')

    def conn(self):
        if not self.connDB:
            self.connDB = ConnDB(self)
        if self.connDB.exec_() == ConnDB.Accepted:
            ip, name = self.connDB.get_output()
            if ip == '10.5.264.8' and name == 'Кравцов Л.С.':
                QtWidgets.QMessageBox.information(QtWidgets.QMessageBox(), 'Подключение к БД',
                                                  'Подключение к базе данных прошло успешно.')
            else:
                QtWidgets.QMessageBox.warning(QtWidgets.QMessageBox(), 'Подключение к БД',
                                              'Не удалось подключиться к базе данных.\n'
                                              'Попробуйте снов или завершите работу.')

    @QtCore.pyqtSlot()
    def open_refresh(self):
        if self.ui.treeWidget.currentItem().text(self.ui.treeWidget.currentColumn()) == 'Программа строительства на год':
            self.loadProjs()

    @QtCore.pyqtSlot()
    def open_about(self):
        if not self.aboutWin:
            self.aboutWin = mainwindow.AboutDialog(self)
        self.aboutWin.exec_()

    @QtCore.pyqtSlot()
    def open_add(self):
        if self.ui.treeWidget.currentItem().text(self.ui.treeWidget.currentColumn()) == 'Программа строительства на год':
            self.addProjs()
        elif self.ui.treeWidget.currentItem().text(self.ui.treeWidget.currentColumn()) == 'Справочник трудовых ресурсов':
            #QtWidgets.QMessageBox.information(QtWidgets.QMessageBox(), 'asdfasdf', 'Добавка туда.')
            pass
        elif self.ui.treeWidget.currentItem().text(self.ui.treeWidget.currentColumn()) == 'Справочник материальных ресурсов':
            self.addMat()

    @QtCore.pyqtSlot()
    def open_edit(self):
        if self.ui.treeWidget.currentItem().text(self.ui.treeWidget.currentColumn()) == 'Программа строительства на год':
            self.editProjs()

    @QtCore.pyqtSlot()
    def open_del(self):
        if self.ui.treeWidget.currentItem().text(self.ui.treeWidget.currentColumn()) == 'Программа строительства на год':
            self.delProjs()
        elif self.ui.treeWidget.currentItem().text(self.ui.treeWidget.currentColumn()) == 'Справочник трудовых ресурсов':
            #QtWidgets.QMessageBox.information(QtWidgets.QMessageBox(), 'asdfasdf', 'Добавка туда.')
            pass
        elif self.ui.treeWidget.currentItem().text(self.ui.treeWidget.currentColumn()) == 'Справочник материальных ресурсов':
            #QtWidgets.QMessageBox.information(QtWidgets.QMessageBox(), 'ahdfgfyjytf', 'Добавка сюда.')
            pass

    # обновление словаря
    def upMatDct(self, d, key, value):
        if key in d:
            d[key].append(value)
        else:
            d[key] = [value]


def loadSession():
    metadata = Base.metadata
    Session = sessionmaker(bind=engine)
    session = Session()
    return session


if __name__ == "__main__":
    import sys

    session = loadSession()
    app = QtWidgets.QApplication(sys.argv)
    window = MainWindow(session)
    window.show()
    sys.exit(app.exec_())

"""
class AboutDialog(QtWidgets.QDialog):
    def __init__(self, parent=None):
        QtWidgets.QDialog.__init__(self, parent)

        self.setFixedWidth(600)
        self.setFixedHeight(250)

        QBtn = QtWidgets.QDialogButtonBox.Ok
        self.buttonBox = QtWidgets.QDialogButtonBox(QBtn)
        self.buttonBox.accepted.connect(self.accept)
        self.buttonBox.rejected.connect(self.reject)

        layout = QtWidgets.QVBoxLayout()

        self.setWindowTitle("О программе")
        title1 = QtWidgets.QLabel("Московский государственный строительный университет")
        font = title1.font()
        font.setPointSize(14)
        title1.setFont(font)

        title2 = QtWidgets.QLabel("Кафедра Информационных систем, технологий и автоматизации в строительстве")
        font = title2.font()
        font.setPointSize(11)
        title2.setFont(font)

        title3 = QtWidgets.QLabel("Демонстрационная версия подсистемы подготовки производства АСУ АО \"СтройФасад\"")
        font = title3.font()
        font.setPointSize(10)
        title3.setFont(font)

        layout.addWidget(title1)
        layout.addWidget(title2)
        layout.addWidget(title3)

        layout.addWidget(QtWidgets.QLabel("v0.01"))
        layout.addWidget(QtWidgets.QLabel("Copyright Леонид Кравцов, курс 4, группа 2"))
        layout.addWidget(QtWidgets.QLabel("2019 г."))
        layout.addWidget(self.buttonBox)

        self.setLayout(layout)
"""